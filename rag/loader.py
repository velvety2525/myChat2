"""규정 폴더의 파일을 읽어 텍스트로 변환한다."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class Document:
    source: str  # 폴더 기준 상대 경로
    text: str
    # (시작 오프셋, 페이지 번호) 목록. PDF만 채워진다.
    page_offsets: list[tuple[int, int]] = field(default_factory=list)
    # 페이지 번호 -> 규정명. 여러 규정을 묶은 규정집 PDF에서 페이지 머리글로 채운다.
    page_titles: dict[int, str] = field(default_factory=dict)

    def page_at(self, offset: int) -> int | None:
        page = None
        for start, num in self.page_offsets:
            if start > offset:
                break
            page = num
        return page


def _read_text_file(path: Path) -> str:
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


# 규정집 페이지 머리글. 예)
#   경북전문대학교
#   취 업 규 칙
#   문서번호 KCS-B-301
#   제ㆍ개정일 2025. 08. 01.
#   개정번호 14
#   페이지 17/46
#   서식B406-05   A4(210×297)
PAGE_HEADER_RE = re.compile(
    r"\A(?P<title>(?:[^\n]*\n){1,4}?)[ \t]*문서번호[ \t]*(?P<docno>[^\n]*)\n(?:[^\n]*\n){0,5}?[ \t]*서식[^\n]*\n"
)


def _normalize_title(lines: list[str]) -> str:
    # 첫 줄이 기관명이고 다음 줄부터 규정명인 형식: 기관명 줄은 버린다
    if len(lines) >= 2:
        lines = lines[1:]
    title = " ".join(lines)
    words = title.split()
    if words and all(len(w) == 1 for w in words):  # "취 업 규 칙" -> "취업규칙"
        title = "".join(words)
    return re.sub(r"\s+", " ", title).strip()


def _split_page_header(text: str) -> tuple[str | None, str]:
    """페이지 머리글이 있으면 (규정명, 머리글을 뺀 본문)을 돌려준다."""
    m = PAGE_HEADER_RE.match(text)
    if not m:
        return None, text
    lines = [ln.strip() for ln in m.group("title").splitlines() if ln.strip()]
    title = _normalize_title(lines)
    docno = m.group("docno").strip()
    if docno:
        title = f"{title}({docno})"
    return title or None, text[m.end():]


def _read_pdf(path: Path) -> tuple[str, list[tuple[int, int]], dict[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    offsets: list[tuple[int, int]] = []
    titles: dict[int, str] = {}
    pos = 0
    for i, page in enumerate(reader.pages, start=1):
        title, text = _split_page_header((page.extract_text() or "").strip() + "\n")
        text = text.strip()
        if title:
            titles[i] = title
        offsets.append((pos, i))
        parts.append(text)
        pos += len(text) + 2  # "\n\n" 구분자
    return "\n\n".join(parts), offsets, titles


def _read_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def load_document(path: Path, root: Path) -> Document | None:
    ext = path.suffix.lower()
    source = str(path.relative_to(root))
    if ext == ".pdf":
        text, offsets, titles = _read_pdf(path)
        return Document(source, text, offsets, titles)
    if ext == ".docx":
        return Document(source, _read_docx(path))
    if ext in (".txt", ".md"):
        return Document(source, _read_text_file(path))
    return None


def list_files(root: Path) -> list[Path]:
    if not root.exists():
        return []
    return sorted(
        p for p in root.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS and not p.name.startswith((".", "~$"))
    )
